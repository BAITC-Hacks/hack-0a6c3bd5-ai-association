"""Ограниченный локальный разбор коротких спецификаций без выполнения содержимого."""

from __future__ import annotations

import io
import re
import warnings as python_warnings
import zipfile
from pathlib import PurePosixPath

from defusedxml import ElementTree

from app.errors import ApiError

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_LINES = 50
MAX_PAGES = 10
MAX_TEXT = 200_000
MAX_XML_BYTES = 8 * 1024 * 1024
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_IMAGE_PIXELS = 16_000_000


def _limit(message: str) -> None:
    raise ApiError(422, "DOCUMENT_LIMIT_EXCEEDED", message)


def _failed(message: str = "Не удалось прочитать документ. Проверьте файл или вставьте строки текстом.") -> None:
    raise ApiError(422, "DOCUMENT_PARSE_FAILED", message)


def _unsupported() -> None:
    raise ApiError(415, "UNSUPPORTED_FILE_TYPE", "Поддерживаются XLSX, DOCX, текстовый PDF и JPEG. Содержимое должно соответствовать расширению.")


def _validate_package(content: bytes, suffix: str) -> None:
    """Проверяем ZIP до библиотек; ничего не извлекаем в файловую систему."""
    required = "xl/workbook.xml" if suffix == ".xlsx" else "word/document.xml"
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        infos = package.infolist()
        names = [item.filename for item in infos]
        if required not in names or "[Content_Types].xml" not in names:
            _unsupported()
        if len(infos) > 2000 or sum(item.file_size for item in infos) > MAX_EXPANDED_BYTES:
            _limit("Распакованный документ слишком большой. Оставьте только короткую спецификацию.")
        if len(names) != len(set(names)):
            _failed("В документе повторяются внутренние части. Сохраните его заново.")
        nodes = 0
        for item in infos:
            path = PurePosixPath(item.filename)
            if path.is_absolute() or ".." in path.parts or "\\" in item.filename or item.flag_bits & 1:
                _failed("Зашифрованный или некорректный пакет документа не поддерживается.")
            if item.file_size > MAX_XML_BYTES:
                _limit("Одна из частей документа слишком велика. Уменьшите файл.")
            if item.file_size > 1024 * 1024 and item.file_size > max(1, item.compress_size) * 200:
                _limit("Документ превышает допустимую сложность распаковки.")
            lowered = item.filename.casefold()
            if lowered.endswith("vbaproject.bin") or "/embeddings/" in lowered:
                _failed("Документы с макросами и вложенными объектами не поддерживаются.")
            if lowered.endswith((".xml", ".rels")):
                root = ElementTree.fromstring(package.read(item))
                for node in root.iter():
                    nodes += 1
                    if nodes > 200_000:
                        _limit("В документе слишком много элементов. Оставьте короткую спецификацию.")
                    if node.tag.endswith("}Relationship") and node.get("TargetMode", "").casefold() == "external":
                        _failed("Документ содержит внешние ссылки. Сохраните копию без внешних связей.")
                    if node.tag.endswith("}Override") and "macroenabled" in node.get("ContentType", "").casefold():
                        _failed("Документы с макросами не поддерживаются.")


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _label(value: str) -> str:
    return re.sub(r"[\s.\-_№/]+", "", value.casefold())


QUERY_HEADERS = {"наименование", "наименованиетовара", "товар", "название", "описание", "product", "description", "name", "item"}
SKU_HEADERS = {"артикул", "код", "кодтовара", "sku", "article", "productcode"}
QUANTITY_HEADERS = {"количество", "колво", "кол", "quantity", "qty", "объем", "объём"}
UNIT_HEADERS = {"едизм", "единица", "единицаизмерения", "ед", "unit", "units", "uom"}
UNITS = {"шт": "шт", "штук": "шт", "штуки": "шт", "штука": "шт", "pcs": "шт", "pc": "шт", "м": "м", "метр": "м", "метров": "м", "метры": "м", "m": "м"}


def _header(row: list[str]) -> dict[str, int] | None:
    columns = {}
    for index, cell in enumerate(row):
        label = _label(cell)
        key = "sku" if label in SKU_HEADERS else "query" if label in QUERY_HEADERS else "quantity" if label in QUANTITY_HEADERS else "unit" if label in UNIT_HEADERS else None
        if key:
            if key in columns:
                _failed("В таблице повторяются заголовки колонок. Оставьте одну колонку каждого типа.")
            columns[key] = index
    if ("query" in columns or "sku" in columns) and ("quantity" in columns or "unit" in columns):
        return columns
    return None


class _Lines:
    def __init__(self):
        self.lines: list[dict] = []
        self.warnings: list[str] = []
        self.text_length = 0

    def add(self, query: str, raw_quantity: str = "", raw_unit: str = "") -> None:
        query = " ".join(query.split())
        if query.startswith("="):
            _failed("Артикул или название задано формулой. Замените формулу текстом.")
        if not query:
            if raw_quantity or raw_unit:
                _failed("В строке с количеством отсутствует название или артикул. Заполните его в исходном файле.")
            return
        self.text_length += len(query) + len(raw_quantity) + len(raw_unit)
        if self.text_length > MAX_TEXT or len(query) > 2000:
            _limit("Текст спецификации слишком длинный. Сократите название или число строк.")
        if len(self.lines) >= MAX_LINES:
            _limit("Поддерживается не более 50 строк спецификации.")
        line_id = str(len(self.lines) + 1)
        quantity = int(raw_quantity) if re.fullmatch(r"[0-9]{1,9}", raw_quantity) and int(raw_quantity) > 0 else None
        unit = UNITS.get(raw_unit.casefold().strip().rstrip("."))
        if quantity is None:
            self.warnings.append(f"Строка {line_id}: количество не определено однозначно. Укажите положительное целое число без формул.")
        if unit is None:
            self.warnings.append(f"Строка {line_id}: единица не определена. Поддерживаются шт и м; проверьте исходный файл.")
        self.lines.append({"line_id": line_id, "query": query, "quantity": quantity, "unit": unit})

    def rows(self, rows) -> None:
        columns = None
        for values in rows:
            row = [_text(value) for value in values]
            while row and not row[-1]:
                row.pop()
            if not any(row):
                continue
            detected = _header(row)
            if detected is not None:
                columns = detected
                continue
            if len(row) == 1:
                self._single(row[0])
                continue
            if columns:
                def cell(key):
                    index = columns.get(key)
                    return row[index] if index is not None and index < len(row) else ""
                sku, description = cell("sku"), cell("query")
                query = " ".join(dict.fromkeys(value for value in (sku, description) if value))
                self.add(query, cell("quantity"), cell("unit"))
            elif len(row) == 3:
                if row[0].startswith("="):
                    _failed("Артикул или название задано формулой. Замените формулу текстом.")
                self.add(*row)
            else:
                self.add(" | ".join(row))

    def plain(self, value: str) -> None:
        self.rows(re.split(r"\s*[|;\t]\s*|\s{2,}", line.strip()) for line in value.splitlines() if line.strip())

    def _single(self, text: str) -> None:
        if re.fullmatch(r"(?:демонстрационная|демо|demo)\s+(?:спецификация|specification)(?:\s+контур)?", text, re.I):
            self.warnings.append("Демонстрационный файл: строки нужно проверить перед созданием предложения.")
            return
        match = re.fullmatch(r"(.+?)\s+(?:[-—:]\s*)?([+-]?[0-9]+(?:[.,][0-9]+)?)\s*(шт\.?|штук|штуки|штука|pcs|pc|м\.?|метр|метров|метры|m)", text, re.I)
        if match:
            self.add(*match.groups())
        else:
            self.add(text)

    def result(self) -> dict:
        if not self.lines:
            _failed("В документе не найдены строки спецификации. Вставьте название, количество и единицу текстом.")
        return {"lines": self.lines, "warnings": list(dict.fromkeys(self.warnings))}


def _xlsx(content: bytes) -> dict:
    from openpyxl import load_workbook

    result = _Lines()
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if len(book.worksheets) > 20:
            _limit("Поддерживается не более 20 листов таблицы.")
        for sheet in book.worksheets:
            if (sheet.max_row or 0) > 10_000 or (sheet.max_column or 0) > 64:
                _limit("Таблица слишком велика: уберите пустое форматирование и лишние колонки.")
            def rows():
                for number, cells in enumerate(sheet.iter_rows(), start=1):
                    if number > 10_000 or len(cells) > 64:
                        _limit("Таблица превышает допустимые размеры.")
                    row = []
                    for cell in cells:
                        value = cell.value
                        if cell.data_type == "f":
                            value = "="
                        elif isinstance(value, (int, float)) and not isinstance(value, bool) and re.fullmatch(r"0{2,}", cell.number_format or ""):
                            value = str(int(value)).zfill(len(cell.number_format)) if int(value) == value else str(value)
                        row.append(value)
                    yield row
            result.rows(rows())
    finally:
        book.close()
    return result.result()


def _docx(content: bytes) -> dict:
    from docx import Document
    from docx.table import Table

    result = _Lines()
    document = Document(io.BytesIO(content))
    for block in document.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                if any(cell.tables for cell in row.cells):
                    _failed("Вложенные таблицы не поддерживаются. Перенесите строки в простую таблицу.")
            result.rows([cell.text for cell in row.cells] for row in block.rows)
        else:
            result.plain(block.text)
    return result.result()


def _pdf(content: bytes) -> dict:
    from pypdf import apply_configuration
    from pypdf.errors import LimitReachedError

    try:
        with apply_configuration(
            maximum_declared_stream_length=MAX_XML_BYTES,
            array_based_stream_maximum_output_length=MAX_XML_BYTES,
            zlib_maximum_output_length=MAX_XML_BYTES,
            lzw_maximum_output_length=MAX_XML_BYTES,
            run_length_maximum_output_length=MAX_XML_BYTES,
            page_tree_maximum_entries=1000,
            page_tree_maximum_depth=30,
            xform_maximum_invocations_per_extraction=100,
        ):
            return _pdf_text(content)
    except LimitReachedError:
        _limit("PDF превышает допустимую сложность. Сохраните короткую текстовую спецификацию.")


def _pdf_text(content: bytes) -> dict:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content), strict=True)
    if reader.is_encrypted:
        _failed("PDF защищён паролем. Сохраните доступную для чтения копию.")
    if len(reader.pages) > MAX_PAGES:
        _limit("Поддерживается не более 10 страниц PDF.")
    result = _Lines()
    text_size = 0
    for page in reader.pages:
        stream = page.get_contents()
        if stream is not None and len(stream.get_data()) > MAX_XML_BYTES:
            _limit("Страница PDF слишком сложная. Сохраните короткую текстовую спецификацию.")
        text = (page.extract_text(extraction_mode="layout") or "") if stream is not None else ""
        text_size += len(text)
        if text_size > MAX_TEXT:
            _limit("PDF содержит слишком много текста.")
        if not text.strip():
            _failed("PDF содержит страницу без текстового слоя. Скан-PDF не поддерживается; загрузите JPEG или вставьте текст.")
        result.plain(text)
    return result.result()


def _jpeg(content: bytes, demo_mode: bool) -> dict:
    from PIL import Image
    from app.llm import extract_image_lines

    with python_warnings.catch_warnings():
        python_warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(content)) as image:
            if image.format != "JPEG":
                _unsupported()
            if image.width * image.height > MAX_IMAGE_PIXELS:
                _limit("JPEG слишком большой после декодирования. Уменьшите его до 16 мегапикселей.")
            image.verify()
        with Image.open(io.BytesIO(content)) as image:
            image.load()
    parsed = extract_image_lines(content, demo_mode=demo_mode)
    if len(parsed["lines"]) > MAX_LINES:
        _limit("Поддерживается не более 50 строк спецификации.")
    return parsed


def parse_document(filename: str, content: bytes, demo_mode: bool = True) -> dict:
    """Возвращает только текстовые строки; название файла никогда не используется как путь."""
    if len(content) > MAX_FILE_BYTES:
        raise ApiError(413, "FILE_TOO_LARGE", "Максимальный размер файла — 10 MiB.")
    suffix = PurePosixPath(filename.replace("\\", "/")).suffix.casefold()
    signatures = {".xlsx": b"PK\x03\x04", ".docx": b"PK\x03\x04", ".pdf": b"%PDF-", ".jpg": b"\xff\xd8\xff", ".jpeg": b"\xff\xd8\xff"}
    if suffix not in signatures or not content.startswith(signatures[suffix]):
        _unsupported()
    try:
        if suffix in {".xlsx", ".docx"}:
            _validate_package(content, suffix)
            return _xlsx(content) if suffix == ".xlsx" else _docx(content)
        if suffix == ".pdf":
            return _pdf(content)
        return _jpeg(content, demo_mode)
    except ApiError:
        raise
    except Exception:
        _failed()
