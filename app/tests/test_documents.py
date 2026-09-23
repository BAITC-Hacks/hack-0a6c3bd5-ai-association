"""Проверки безопасного разбора файла и неоднозначностей количества."""

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from app.documents import MAX_FILE_BYTES, parse_document
from app.errors import ApiError


def xlsx(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def docx(rows=None, paragraphs=None):
    document = Document()
    for text in paragraphs or []:
        document.add_paragraph(text)
    if rows:
        table = document.add_table(rows=0, cols=len(rows[0]))
        for values in rows:
            for cell, value in zip(table.add_row().cells, values):
                cell.text = str(value)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf(text=None, pages=1):
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=600, height=800)
        if text is not None:
            font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
            page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
            stream = DecodedStreamObject()
            stream.set_data(('BT /F1 12 Tf 40 760 Td (' + text + ') Tj ET').encode('ascii'))
            page[NameObject('/Contents')] = writer._add_object(stream)
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def rewrite_zip(content, replacements):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as source, zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as dest:
        for entry in source.infolist():
            dest.writestr(entry.filename, replacements.pop(entry.filename, source.read(entry)))
        for name, data in replacements.items():
            dest.writestr(name, data)
    return output.getvalue()


def failure(filename, content, code):
    with pytest.raises(ApiError) as error:
        parse_document(filename, content, True)
    assert error.value.code == code
    return error.value


@pytest.mark.parametrize('extension,builder', [('xlsx', xlsx), ('docx', docx)])
def test_table_headers_and_leading_zero(extension, builder):
    parsed = parse_document('spec.' + extension, builder([['Количество', 'Ед. изм.', 'Артикул', 'Наименование'], ['2', 'шт.', '000123', 'Автомат 160 А']]), True)
    assert parsed == {'lines': [{'line_id': '1', 'query': '000123 Автомат 160 А', 'quantity': 2, 'unit': 'шт'}], 'warnings': []}


def test_xlsx_numeric_sku_display_zeros():
    book = Workbook()
    book.active.append(['Артикул', 'Количество', 'Ед.изм.'])
    book.active.append([123, 2, 'шт'])
    book.active['A2'].number_format = '000000'
    stream = io.BytesIO()
    book.save(stream)
    assert parse_document('spec.xlsx', stream.getvalue(), True)['lines'][0]['query'] == '000123'


def test_quantity_formulas_and_fractions_require_correction():
    parsed = parse_document('spec.xlsx', xlsx([['SKU', 'Qty', 'Unit'], ['one', '=1+1', 'pcs'], ['two', 1.5, 'м'], ['three', '-2', 'кг']]), True)
    assert [row['quantity'] for row in parsed['lines']] == [None, None, None]
    assert parsed['lines'][2]['unit'] is None
    assert len(parsed['warnings']) == 4


def test_query_formula_is_not_exposed_as_product():
    failure('spec.xlsx', xlsx([['SKU', 'Qty', 'Unit'], ['="test"', 2, 'pcs']]), 'DOCUMENT_PARSE_FAILED')


def test_no_quantity_guessed_from_current_or_sku():
    parsed = parse_document('spec.docx', docx(paragraphs=['Автомат 160 А', '200300285_', 'Кабель 3x2.5 | 5 | м']), True)
    assert parsed['lines'][0]['quantity'] is None
    assert parsed['lines'][1]['quantity'] is None
    assert parsed['lines'][2] == {'line_id': '3', 'query': 'Кабель 3x2.5', 'quantity': 5, 'unit': 'м'}


def test_docx_preserves_paragraph_and_table_order():
    document = Document()
    document.add_paragraph('first | 2 | шт')
    table = document.add_table(rows=1, cols=3)
    for cell, text in zip(table.rows[0].cells, ['second', '3', 'м']):
        cell.text = text
    document.add_paragraph('third 4 шт')
    stream = io.BytesIO()
    document.save(stream)
    assert [line['query'] for line in parse_document('spec.docx', stream.getvalue(), True)['lines']] == ['first', 'second', 'third']


def test_text_pdf():
    assert parse_document('spec.pdf', pdf('DEMO-160-AVAILABLE | 2 | pcs'), True)['lines'][0] == {'line_id': '1', 'query': 'DEMO-160-AVAILABLE', 'quantity': 2, 'unit': 'шт'}


def test_pdf_page_and_scan_limits():
    failure('spec.pdf', pdf('SKU | 1 | pcs', pages=11), 'DOCUMENT_LIMIT_EXCEEDED')
    error = failure('scan.pdf', pdf(), 'DOCUMENT_PARSE_FAILED')
    assert 'Скан-PDF' in error.message


@pytest.mark.parametrize('extension,builder', [('xlsx', xlsx), ('docx', docx)])
def test_line_limit_never_truncates(extension, builder):
    failure('spec.' + extension, builder([['sku', 1, 'шт']] * 51), 'DOCUMENT_LIMIT_EXCEEDED')


def test_empty_table_and_oversized_query():
    failure('spec.xlsx', xlsx([['SKU', 'Qty', 'Unit']]), 'DOCUMENT_PARSE_FAILED')
    failure('spec.docx', docx(paragraphs=['a' * 2001]), 'DOCUMENT_LIMIT_EXCEEDED')


@pytest.mark.parametrize('filename,content', [('spec.xls', b'fake'), ('spec.jpg', b'%PDF-1.4'), ('spec.docx', xlsx([['sku', 2, 'pcs']])), ('spec.pdf', b'not pdf')])
def test_extension_is_not_a_type_proof(filename, content):
    failure(filename, content, 'UNSUPPORTED_FILE_TYPE')


def test_size_and_corrupt_package():
    assert failure('spec.xlsx', b'x' * (MAX_FILE_BYTES + 1), 'FILE_TOO_LARGE').status == 413
    failure('spec.xlsx', b'PK\x03\x04corrupt', 'DOCUMENT_PARSE_FAILED')


def test_ooxml_external_links_and_entities_rejected():
    original = xlsx([['sku', 2, 'pcs']])
    external = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.com/" TargetMode="External"/></Relationships>'
    failure('spec.xlsx', rewrite_zip(original, {'xl/worksheets/_rels/sheet1.xml.rels': external}), 'DOCUMENT_PARSE_FAILED')
    entity = b'<!DOCTYPE x [<!ENTITY p SYSTEM "file:///C:/secret">]><x>&p;</x>'
    failure('spec.xlsx', rewrite_zip(original, {'extra.xml': entity}), 'DOCUMENT_PARSE_FAILED')


def test_expanded_zip_limit():
    original = xlsx([['sku', 2, 'pcs']])
    bomb = b'<x>' + b' ' * (2 * 1024 * 1024) + b'</x>'
    failure('spec.xlsx', rewrite_zip(original, {'extra.xml': bomb}), 'DOCUMENT_LIMIT_EXCEEDED')


def test_jpeg_decoding_and_ocr_hook(monkeypatch):
    stream = io.BytesIO()
    Image.new('RGB', (100, 50), 'white').save(stream, 'JPEG')
    expected = {'lines': [{'line_id': '1', 'query': 'sku', 'quantity': 2, 'unit': 'шт'}], 'warnings': []}
    monkeypatch.setattr('app.llm.extract_image_lines', lambda content, demo_mode: expected, raising=False)
    assert parse_document('sample.JPEG', stream.getvalue(), True) == expected
    failure('sample.jpg', b'\xff\xd8\xffbroken', 'DOCUMENT_PARSE_FAILED')


def test_committed_samples():
    directory = Path(__file__).resolve().parents[2] / 'fixtures' / 'uploads'
    for extension in ('xlsx', 'docx', 'pdf'):
        path = directory / ('sample.' + extension)
        parsed = parse_document(path.name, path.read_bytes(), True)
        assert [(row['query'], row['quantity'], row['unit']) for row in parsed['lines']] == [('DEMO-160-AVAILABLE', 2, 'шт'), ('200300285_', 1, 'шт')]


def test_sku_does_not_erase_incompatible_name_requirement():
    parsed = parse_document('spec.xlsx', xlsx([['Артикул', 'Наименование', 'Количество', 'Ед.изм.'], ['DEMO-160-AVAILABLE', 'Автомат 250 А', 2, 'шт']]), True)
    assert parsed['lines'][0]['query'] == 'DEMO-160-AVAILABLE Автомат 250 А'


def test_pdf_compressed_content_limit():
    from pypdf.generic import DecodedStreamObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    stream = DecodedStreamObject()
    stream.set_data(b' ' * (9 * 1024 * 1024))
    page[NameObject('/Contents')] = writer._add_object(stream.flate_encode())
    output = io.BytesIO()
    writer.write(output)
    failure('oversized.pdf', output.getvalue(), 'DOCUMENT_LIMIT_EXCEEDED')
